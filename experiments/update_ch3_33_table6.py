#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""表6（3.3）只保留 GM(S)：删除「有效候选率」列，数值改在正文中给出。

理由
----
全文其余各表都以比值型指标（GM(S)、GM(G_cross)、GM(G_mut)、GM(F)）为主，
有效候选率在表6 中并未提供额外判别力（FULL 87.70% 低于 A6/B3/B4），
反而让主指标不够突出。因此从表中删除该列。

注意：有效候选率承载的"两个组件存在叠加效应"这一结论予以保留——
A7 为 86.47%、低于仅关闭单一组件的 A6 与 B3，该论述在分析段中以文字给出，
不进表，因此删除列不会丢掉这条结论。

用法
----
    python experiments/update_ch3_33_table6.py
"""
import shutil
import sys
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from experiments.docx_util import delete_column, find_paragraph, set_text  # noqa: E402

DOC = ROOT / "doc" / "第三章_修订版_四节实验数据整理_表4拆分.docx"
BACKUP = ROOT / "doc" / "第三章_修订版_四节实验数据整理_表4拆分_表6删列前备份.docx"

COL = "有效候选率"

# 表6 的识别：表头同时含 Profiling 与 RAG 方式（避免与 3.6 的表11 混淆）
def find_table6(d):
    for i, t in enumerate(d.tables):
        hdr = [c.text.strip() for c in t.rows[0].cells]
        if any("Profiling" in h for h in hdr) and any("RAG" in h for h in hdr):
            return i, t
    raise LookupError("未找到表6")


A_MAIN = "主表统一使用同一批 15 个代表性算子"
OLD_HEAD = "报告有效候选率与最终加速比 GM(S)"
NEW_HEAD = "以最终加速比 GM(S) 为主指标，有效候选率的数值在结果分析中一并给出"


def main():
    if not BACKUP.exists():
        shutil.copy2(DOC, BACKUP)
        print(f"已备份: {BACKUP.name}")
    else:
        print(f"备份已存在，跳过: {BACKUP.name}")

    d = Document(DOC)
    ti, tbl = find_table6(d)
    hdr = [c.text.strip() for c in tbl.rows[0].cells]

    if COL in hdr:
        delete_column(tbl, hdr.index(COL))
        print(f"  - 表6(表#{ti}) 已删除「{COL}」列，剩余列: "
              + " | ".join(c.text.strip() for c in tbl.rows[0].cells))
    else:
        print(f"  - 表6 已无「{COL}」列，跳过")

    _, p = find_paragraph(d, A_MAIN)
    if OLD_HEAD in p.text:
        set_text(p, p.text.strip().replace(OLD_HEAD, NEW_HEAD))
        print("  - 已同步修改主表口径说明段")
    else:
        print("  - 主表口径说明段无需修改")

    d.save(DOC)

    d2 = Document(DOC)
    _, t2 = find_table6(d2)
    print(f"\n已写回: {DOC.name}  (表格总数 {len(d2.tables)})")
    print("表6 内容：")
    for r in t2.rows:
        print("   " + " | ".join(c.text.strip().rjust(14) for c in r.cells))


if __name__ == "__main__":
    main()

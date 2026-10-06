#!/usr/bin/env python3
"""生成 3.4 进化策略消融的 Word 结果表（表7 / 表8 / 表9）。

用法：
    cd /workspace/Agent && python experiments/gen_table789_docx.py

数据源：experiments/results/table{7,8,9}_*.json（由 analyze_logs.py 3.4.x 产出）
输出：
    doc/表7_3.4选择策略消融.docx
    doc/表8_3.4交叉策略消融.docx
    doc/表9_3.4变异策略消融.docx

若某张表的数据尚未产出（n_kernels 全为 0），则跳过该表并提示，
以免生成无意义的空表。
"""
import json
from pathlib import Path
from docx import Document
from docx.shared import Pt

ROOT = Path(__file__).resolve().parent.parent
RES = ROOT / 'experiments/results'


def fmt(v, nd=4):
    """None / 空值统一显示为 —"""
    if v is None or v == '':
        return '—'
    if isinstance(v, float):
        return f'{v:.{nd}f}'
    return str(v)


def pct(v, nd=2):
    return '—' if v is None else f'{v * 100:.{nd}f}%'


SPECS = [
    {
        'src': RES / 'table7_selection.json',
        'out': ROOT / 'doc/表7_3.4选择策略消融.docx',
        'title': '表7 选择策略消融实验结果（3.4.1）',
        'desc': (
            '实验设置：在完整方法基础上仅替换父代选择策略，其余组件保持不变，'
            '比较轮盘赌选择（按适应度归一化为采样概率）与锦标赛选择'
            '（每次随机抽取 k 个候选并取最优者）。'
            '指标上，被选中父代的平均适应度反映选择压力：'
            '该值越高说明选择越偏向当前高适应度个体，种群多样性相应越低。'
        ),
        'cfg_desc': {
            'sel_roulette': '轮盘赌选择（按适应度概率采样）',
            'sel_tournament': '锦标赛选择（k 个候选竞争取优）',
            'sel_uniform': '均匀随机选择（无选择压力，下界基线）',
            'sel_ucb': 'UCB 探索—利用平衡选择（本文提出）',
        },
        'hdr': ['配置', '选择策略', '选择次数', '被选父代平均适应度', 'GM(S)', '参与 Kernel 数'],
        'row': lambda r: [
            r['config'],
            r.get('selection_mode', ''),
            fmt(r.get('n_selections'), 0),
            fmt(r.get('mean_parent_fitness')),
            fmt(r.get('gm_s')),
            fmt(r.get('n_kernels'), 0),
        ],
        'note': '注：被选父代平均适应度 = 所有被选中父代个体适应度的均值，用于衡量选择压力；GM(S) 为最终加速比的几何平均。',
    },
    {
        'src': RES / 'table8_crossover.json',
        'out': ROOT / 'doc/表8_3.4交叉策略消融.docx',
        'title': '表8 交叉策略消融实验结果（3.4.2）',
        'desc': (
            '实验设置：比较两种交叉算子。无约束交叉允许在两个父代代码的任意位置'
            '交换代码片段；受保护交叉则在交换时保留 Kernel 的关键结构（如函数签名、'
            '索引计算与内存访问模式），避免破坏语义。'
            '有效率指交叉产生的后代中通过编译与功能验证的比例；'
            '改进率指优于其较优父代的后代占有效后代的比例。'
        ),
        'cfg_desc': {
            'cross_unconstrained': '无约束交叉（任意片段交换）',
            'cross_protected': '受保护交叉（保留关键结构）',
        },
        'hdr': ['配置', '交叉策略', '交叉次数', '有效后代', '有效率', '优于父代数', '改进率', 'GM(增益)', 'GM(S)'],
        'row': lambda r: [
            r['config'],
            '',
            fmt(r.get('n_crossover'), 0),
            fmt(r.get('n_valid'), 0),
            pct(r.get('valid_rate')),
            fmt(r.get('n_better'), 0),
            pct(r.get('better_ratio')),
            fmt(r.get('gm_g_cross')),
            fmt(r.get('gm_s')),
        ],
        'note': '注：有效率 = 有效后代 / 交叉次数；改进率 = 优于较优父代的后代 / 有效后代；GM(增益) 为交叉操作带来适应度提升的几何平均。',
    },
    {
        'src': RES / 'table9_mutation_docx.json',
        'out': ROOT / 'doc/表9_3.4变异策略消融.docx',
        'title': '表9 变异策略消融实验结果（3.4.3）',
        'desc': (
            '实验设置：比较三种变异强度策略。自适应变异根据当前代数与个体适应度'
            '动态调整改动幅度；均匀变异采用固定强度的局部改写；'
            '激进变异则对算子实现做较彻底的重构。'
            '成功率指变异后适应度优于变异前的个体占有效变异个体的比例。'
        ),
        'cfg_desc': {
            'mut_adaptive': '自适应变异（按代数/适应度调幅）',
            'mut_uniform': '均匀变异（固定幅度局部改写）',
            'mut_aggressive': '激进变异（大幅重构）',
        },
        'hdr': ['配置', '变异策略', '变异次数', '有效个体', '有效率', '改进个体数', '成功率', 'GM(增益)', 'GM(S)'],
        'row': lambda r: [
            r['config'],
            '',
            fmt(r.get('n_mutation'), 0),
            fmt(r.get('n_valid'), 0),
            pct(r.get('valid_rate')),
            fmt(r.get('n_better'), 0),
            pct(r.get('success_rate')),
            fmt(r.get('gm_g_mut')),
            fmt(r.get('gm_s')),
        ],
        'note': '注：有效率 = 有效变异个体 / 变异次数；成功率 = 适应度优于变异前的个体 / 有效变异个体；GM(增益) 为变异操作带来适应度提升的几何平均。',
    },
]

# 表9 的源文件名与其他两张不同，这里统一修正
SPECS[2]['src'] = RES / 'table9_mutation.json'


def build(spec):
    src = spec['src']
    if not src.exists():
        print(f"  跳过 {spec['out'].name}：源文件不存在 {src}")
        return False
    data = json.load(open(src, encoding='utf-8'))
    # 数据未产出（全部 n_kernels=0 / n_for_gm=0）时不生成空表
    ready = any(
        (r.get('n_kernels') or 0) > 0 or (r.get('n_for_gm') or 0) > 0
        for r in data
    )
    if not ready:
        print(f"  跳过 {spec['out'].name}：3.4 数据尚未产出（各配置样本数为 0）")
        return False

    doc = Document()
    doc.add_heading(spec['title'], level=2)

    p = doc.add_paragraph(spec['desc'])
    p.runs[0].font.size = Pt(9)

    hdr = spec['hdr']
    t = doc.add_table(rows=1, cols=len(hdr))
    t.style = 'Table Grid'
    for i, h in enumerate(hdr):
        t.rows[0].cells[i].text = h

    for r in data:
        cells = t.add_row().cells
        vals = spec['row'](r)
        # 第二列填入配置说明
        vals[1] = spec['cfg_desc'].get(r['config'], r.get('selection_mode', ''))
        for i, v in enumerate(vals):
            cells[i].text = v

    doc.add_paragraph()
    fn = doc.add_paragraph(spec['note'])
    fn.runs[0].font.size = Pt(9)

    spec['out'].parent.mkdir(parents=True, exist_ok=True)
    doc.save(spec['out'])
    print(f"  已生成: {spec['out']}")
    return True


if __name__ == '__main__':
    print('生成 3.4 进化策略消融表（表7 / 表8 / 表9）:')
    for s in SPECS:
        build(s)

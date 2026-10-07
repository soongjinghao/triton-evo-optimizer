#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""删除 3.3 中的端到端「表6」，只保留第 0 代候选质量表（原表6b → 改名为表6）。

删除理由
--------
表6 的端到端 GM(S) 在本实验规模下无法区分五个配置（FULL 2.1550 并不高于
A6 2.2566 / B3 2.2715），且该指标经过两代进化与精英保留后已被稀释。
3.3 的结论改由第 0 代候选质量（组件真正起作用的层次）承载。

删除范围
--------
  1. 表题：「表6 Profiling 证据与知识库检索策略消融」
  2. 表体：紧随其后的 4 列表
  3. 分析段：「表6显示，在统一的 15 个算子上……」
  4. 口径说明段改写：原段通篇描述表6 的口径（GM(S) / 有效候选率），
     表删除后会悬空，改为描述本节唯一保留的第 0 代表口径。

改名
----
删除后把正文中的「表6b」统一改为「表6」，避免出现「有 b 而无主表」的
孤儿编号，并使表5 → 表6 → 表7a → 表7 的序列连续。

用法
----
    python experiments/remove_ch3_33_table6.py
"""
import shutil
import sys
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from experiments.docx_util import (                     # noqa: E402
    find_paragraph, remove_element, set_text, table_after,
)

DOC = ROOT / "doc" / "第三章_修订版_四节实验数据整理_表4拆分.docx"
BACKUP = ROOT / "doc" / "第三章_修订版_四节实验数据整理_表4拆分_删表6前备份.docx"

A_CAPTION = "表6 Profiling 证据与知识库检索策略消融"
A_ANALYSIS = "表6显示，在统一的 15 个算子上"
A_SPEC = "主表统一使用同一批 15 个代表性算子"

NEW_SPEC = (
    "本节统一使用上述 15 个算子的子集与正式 r0 日志。"
    "鉴于端到端加速比在本实验规模下无法区分各配置（详见本节结果分析），"
    "本节的度量下沉到 Profiling 证据与知识库检索真正起作用的层次，即第 0 代候选，"
    "采用候选有效率与相对最快种子的加速比 S0 = T_seed / T_candidate 作为指标，"
    "其中候选延迟均经多次测量取中位数以抑制短 Kernel 的计时抖动。"
)


def main():
    if not BACKUP.exists():
        shutil.copy2(DOC, BACKUP)
        print(f"已备份: {BACKUP.name}")
    else:
        print(f"备份已存在，跳过: {BACKUP.name}")

    d = Document(DOC)

    # ---- 删除表6 的表题、表体与分析段 ----
    removed = False
    _, cap = find_paragraph(d, A_CAPTION, required=False)
    if cap is not None:
        tbl = table_after(d, cap)
        if tbl is not None:
            remove_element(tbl)
            print("  - 已删除表6 表体")
        remove_element(cap)
        print("  - 已删除表6 表题")
        removed = True
    else:
        print("  - 表6 表题已不存在，跳过")

    _, p_an = find_paragraph(d, A_ANALYSIS, required=False)
    if p_an is not None:
        remove_element(p_an)
        print("  - 已删除表6 分析段")
        removed = True

    # ---- 口径说明段改写 ----
    _, p_spec = find_paragraph(d, A_SPEC, required=False)
    if p_spec is not None:
        set_text(p_spec, NEW_SPEC)
        print("  - 口径说明段已改写为第 0 代表口径")
    else:
        print("  - 口径说明段已为新版本，跳过")

    # ---- 改名：表6b → 表6 ----
    n = 0
    for p in d.paragraphs:
        if "表6b" in p.text:
            set_text(p, p.text.replace("表6b", "表6"))
            n += 1
    print(f"  - 已将 {n} 处「表6b」改名为「表6」")

    d.save(DOC)

    # ---- 残留检查 ----
    d2 = Document(DOC)
    ps = [p.text.strip() for p in d2.paragraphs]
    leftover = [(i, t[:60]) for i, t in enumerate(ps) if "表6b" in t]
    print(f"\n已写回: {DOC.name}  (段落 {len(ps)}，表格 {len(d2.tables)})")
    if leftover:
        print(f"⚠ 仍有 {len(leftover)} 处「表6b」: {leftover}")
    else:
        print("  ✓ 全文已无「表6b」；表6 现指第 0 代候选质量表")
    print("  幂等，可重复运行")


if __name__ == "__main__":
    main()
